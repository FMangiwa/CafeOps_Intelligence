from __future__ import annotations
import argparse, json, os, sys, time
from pathlib import Path
from typing import Any
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

def clean(v: Any):
    if hasattr(v, "isoformat"): return v.isoformat()
    if isinstance(v, str):
        v=v.strip()
        return v if v else None
    return v

def load_config(path: Path):
    cfg=json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(cfg.get("sheets"),list): raise ValueError("Mapping must contain a 'sheets' list.")
    return cfg

def records_from_sheet(ws, mapping):
    rows=ws.iter_rows(values_only=True)
    try: headers=[str(x).strip() if x is not None else "" for x in next(rows)]
    except StopIteration: return [], ["Worksheet is empty"]
    required=mapping.get("required_headers", list(mapping.get("columns",{}).keys()))
    missing=[h for h in required if h not in headers]
    if missing: return [], [f"Missing required headers: {missing}"]
    colmap=mapping.get("columns",{}); defaults=mapping.get("defaults",{})
    records=[]; errors=[]
    for excel_row,row in enumerate(rows,start=2):
        if mapping.get("skip_empty_rows",True) and all(x is None or str(x).strip()=="" for x in row): continue
        obj={}
        for src,dst in colmap.items():
            idx=headers.index(src)
            obj[dst]=clean(row[idx] if idx<len(row) else None)
        for k,v in defaults.items():
            if k=="metric_name_from_key":
                key=obj.get("metric_key")
                if key: obj["metric_name"]=str(key).replace("_"," ").title()
            elif k not in obj or obj[k] is None:
                obj[k]=v
        if not any(v is not None for v in obj.values()):
            errors.append(f"Row {excel_row}: mapped fields are all empty"); continue
        records.append((excel_row,obj))
    return records,errors

def main():
    ap=argparse.ArgumentParser(description="Validate and optionally seed Excel rows into Supabase.")
    ap.add_argument("--workbook",required=True)
    ap.add_argument("--mapping",default="config/seed_mapping.json")
    ap.add_argument("--apply",action="store_true",help="Perform database writes; default is dry-run.")
    args=ap.parse_args(); load_dotenv()
    book=Path(args.workbook); mapfile=Path(args.mapping)
    if not book.exists(): print(f"ERROR: workbook not found: {book}"); return 2
    if not mapfile.exists(): print(f"ERROR: mapping not found: {mapfile}"); return 2
    try: cfg=load_config(mapfile)
    except Exception as e: print(f"ERROR: invalid mapping: {e}"); return 2
    if not cfg["sheets"]: print("ERROR: no sheet mappings configured.",mapfile); return 2
    wb=load_workbook(book,read_only=True,data_only=True)
    base=os.getenv("SUPABASE_URL","").rstrip("/"); key=os.getenv("SUPABASE_SERVICE_ROLE_KEY","")
    if args.apply and (not base.startswith("https://") or not key or key=="replace_me"):
        print("ERROR: --apply requires valid SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY in .env"); return 2
    if args.apply and key.startswith("sb_publishable_"):
        print("ERROR: SUPABASE_SERVICE_ROLE_KEY is a publishable key. Use a server-side secret/service-role key; never expose it."); return 2
    session=requests.Session()
    session.headers.update({"apikey":key,"Authorization":f"Bearer {key}","Content-Type":"application/json","Prefer":"return=minimal"})
    batch=int(cfg.get("batch_size",250)); total=0; failed=0
    for m in cfg["sheets"]:
        name=m.get("sheet"); table=m.get("table"); mode=m.get("mode","upsert")
        if name not in wb.sheetnames: print(f"[{name}] ERROR: worksheet not found"); failed+=1; continue
        if not table or mode not in ("insert","upsert"): print(f"[{name}] ERROR: invalid table/mode"); failed+=1; continue
        if mode=="upsert" and not m.get("conflict_columns"): print(f"[{name}] ERROR: upsert requires conflict_columns"); failed+=1; continue
        records,errors=records_from_sheet(wb[name],m)
        for e in errors: print(f"[{name}] ERROR: {e}")
        if errors: failed+=len(errors)
        payload=[r for _,r in records]
        print(f"[{name}] {len(payload)} mapped rows → public.{table} ({mode})")
        total+=len(payload)
        if not args.apply: continue
        params={}
        if mode=="upsert":
            params["on_conflict"]=",".join(m["conflict_columns"])
            session.headers["Prefer"]="resolution=merge-duplicates,return=minimal"
        else: session.headers["Prefer"]="return=minimal"
        for start in range(0,len(payload),batch):
            chunk=payload[start:start+batch]
            try:
                resp=session.post(f"{base}/rest/v1/{table}",params=params,json=chunk,timeout=60)
                if not resp.ok:
                    failed+=len(chunk); print(f"[{name}] API ERROR batch {start//batch+1}: HTTP {resp.status_code}: {resp.text[:1200]}")
                else: print(f"[{name}] inserted/upserted batch {start//batch+1} ({len(chunk)} rows)")
            except requests.RequestException as e:
                failed+=len(chunk); print(f"[{name}] NETWORK ERROR: {e}")
            time.sleep(0.1)
    wb.close()
    print(f"\nSUMMARY: mapped_rows={total}; errors_or_failed_rows={failed}; mode={'APPLY' if args.apply else 'DRY-RUN'}")
    if not args.apply: print("DRY-RUN ONLY: no database writes were made.")
    return 1 if failed else 0

if __name__=="__main__": sys.exit(main())
